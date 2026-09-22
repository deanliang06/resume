from __future__ import annotations

import copy
import hashlib
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .errors import InvalidInput, UnsupportedTemplate
from .models import ParsedResume, Project, TailoredResume

SECTION_NAMES = ("Education:", "Experience:", "Projects:", "Skills:")
SKILL_LABELS = ("Languages", "Libraries", "Web & Database", "Tools/Infra", "Hobbies/Other")
URL_RE = re.compile(r"https?://[^\s)]+/?")


@dataclass(frozen=True)
class TemplateMap:
    section_indexes: dict[str, int]
    project_heading_indexes: tuple[int, ...]
    project_blocks: tuple[tuple[int, int], ...]
    skill_indexes: dict[str, int]

    @property
    def editable_indexes(self) -> set[int]:
        indexes = set(self.skill_indexes.values()) - {self.skill_indexes["Hobbies/Other"]}
        for start, end in self.project_blocks:
            indexes.update(range(start, end))
        return indexes


def validate_docx_package(path: Path, max_expanded_bytes: int) -> None:
    if not zipfile.is_zipfile(path):
        raise InvalidInput("The upload is not a valid DOCX package.")
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
            if total > max_expanded_bytes:
                raise InvalidInput("The expanded DOCX package is too large.")
            if "word/document.xml" not in archive.namelist():
                raise InvalidInput("The DOCX has no main document part.")
            for item in archive.infolist():
                normalized = item.filename.replace("\\", "/")
                if normalized.startswith("/") or "../" in normalized:
                    raise InvalidInput("The DOCX contains an unsafe package path.")
    except (OSError, zipfile.BadZipFile) as exc:
        raise InvalidInput("The DOCX package could not be read.") from exc


def map_template(document: DocumentType) -> TemplateMap:
    texts = [p.text.strip() for p in document.paragraphs]
    section_indexes: dict[str, int] = {}
    for name in SECTION_NAMES:
        matches = [i for i, text in enumerate(texts) if text == name]
        if len(matches) != 1:
            raise UnsupportedTemplate(f"Expected exactly one {name} section heading.")
        section_indexes[name] = matches[0]
    ordered = [section_indexes[name] for name in SECTION_NAMES]
    if ordered != sorted(ordered):
        raise UnsupportedTemplate("Resume sections are not in the expected order.")

    project_start = section_indexes["Projects:"] + 1
    skills_start = section_indexes["Skills:"]
    heading_indexes = []
    for i in range(project_start, skills_start):
        p = document.paragraphs[i]
        if p.text.strip() and p.style.name != "List Paragraph":
            heading_indexes.append(i)
    if len(heading_indexes) < 3:
        raise UnsupportedTemplate("At least three project blocks are required.")
    blocks = []
    for position, start in enumerate(heading_indexes):
        end = heading_indexes[position + 1] if position + 1 < len(heading_indexes) else skills_start
        if not any(document.paragraphs[i].style.name == "List Paragraph" and document.paragraphs[i].text.strip() for i in range(start + 1, end)):
            raise UnsupportedTemplate("Each project must contain at least one list bullet.")
        blocks.append((start, end))

    skill_indexes: dict[str, int] = {}
    for label in SKILL_LABELS:
        matches = [i for i in range(skills_start + 1, len(texts)) if texts[i].startswith(label + ":")]
        if len(matches) != 1:
            raise UnsupportedTemplate(f"Expected exactly one {label} skill row.")
        skill_indexes[label] = matches[0]
    if list(skill_indexes.values()) != sorted(skill_indexes.values()):
        raise UnsupportedTemplate("Skill rows are not in the expected order.")
    return TemplateMap(section_indexes, tuple(heading_indexes), tuple(blocks), skill_indexes)


def parse_resume(path: Path) -> tuple[DocumentType, TemplateMap, ParsedResume]:
    try:
        document = Document(path)
    except Exception as exc:
        raise InvalidInput("The DOCX could not be opened.") from exc
    mapping = map_template(document)
    projects = []
    for order, (start, end) in enumerate(mapping.project_blocks):
        heading = document.paragraphs[start].text.strip()
        match = URL_RE.search(heading)
        url = match.group(0) if match else None
        title = heading[: match.start()].rstrip(" (") if match else heading
        bullets = [document.paragraphs[i].text.strip() for i in range(start + 1, end) if document.paragraphs[i].text.strip()]
        projects.append(Project(f"resume-project-{order + 1}", title, url, bullets, order))
    skills = {}
    for label, index in mapping.skill_indexes.items():
        _, _, values = document.paragraphs[index].text.partition(":")
        skills[label] = [item.strip() for item in values.split(",") if item.strip()]
    return document, mapping, ParsedResume(projects, skills, [p.text for p in document.paragraphs])


def protected_digest(document: DocumentType, mapping: TemplateMap) -> str:
    digest = hashlib.sha256()
    for index, paragraph in enumerate(document.paragraphs):
        if index not in mapping.editable_indexes:
            digest.update(paragraph._p.xml.encode("utf-8"))
    return digest.hexdigest()


def _remove_all_runs(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            paragraph._p.remove(child)


def _copy_run_format(source, target) -> None:
    if source is not None and source._r.rPr is not None:
        target._r.insert(0, copy.deepcopy(source._r.rPr))


def _add_hyperlink(paragraph: Paragraph, text: str, url: str, template_run=None, template_rpr=None) -> None:
    relation_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    if template_rpr is not None:
        run.append(copy.deepcopy(template_rpr))
    elif template_run is not None and template_run._r.rPr is not None:
        run.append(copy.deepcopy(template_run._r.rPr))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _set_heading(paragraph: Paragraph, title: str, url: str | None) -> None:
    template_runs = paragraph.runs
    title_template = template_runs[0] if template_runs else None
    link_template = template_runs[-1] if template_runs else title_template
    hyperlink_rpr = None
    for hyperlink in paragraph._p.findall(qn("w:hyperlink")):
        hyperlink_run = hyperlink.find(qn("w:r"))
        if hyperlink_run is not None:
            hyperlink_rpr = hyperlink_run.find(qn("w:rPr"))
            if hyperlink_rpr is not None:
                break
    _remove_all_runs(paragraph)
    title_run = paragraph.add_run(title)
    _copy_run_format(title_template, title_run)
    if url:
        spacer = paragraph.add_run(" (")
        _copy_run_format(title_template, spacer)
        _add_hyperlink(paragraph, url, url, link_template, hyperlink_rpr)
        closer = paragraph.add_run(")")
        _copy_run_format(title_template, closer)


def _set_paragraph_text(paragraph: Paragraph, text: str, *, label: str | None = None) -> None:
    template_runs = paragraph.runs
    normal_template = template_runs[-1] if template_runs else None
    label_template = template_runs[0] if template_runs else normal_template
    _remove_all_runs(paragraph)
    if label is None:
        run = paragraph.add_run(text)
        _copy_run_format(normal_template, run)
        return
    run = paragraph.add_run(label + ":")
    _copy_run_format(label_template, run)
    value = paragraph.add_run(" " + text)
    _copy_run_format(normal_template, value)


def _delete_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def _clone_paragraph_after(paragraph: Paragraph) -> Paragraph:
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    return Paragraph(clone, paragraph._parent)


def apply_tailoring(source: Path, output: Path, tailored: TailoredResume) -> None:
    document, mapping, _ = parse_resume(source)
    before = protected_digest(document, mapping)
    if len(tailored.projects) not in (3, 4):
        raise UnsupportedTemplate("A candidate must contain three or four projects.")
    if len(tailored.projects) > len(mapping.project_blocks):
        raise UnsupportedTemplate("The source template has too few project slots.")

    # Work backwards so removing unused blocks does not invalidate earlier indexes.
    for block_position in range(len(mapping.project_blocks) - 1, -1, -1):
        start, end = mapping.project_blocks[block_position]
        paragraphs = document.paragraphs
        if block_position >= len(tailored.projects):
            for i in range(end - 1, start - 1, -1):
                _delete_paragraph(paragraphs[i])
            continue
        project = tailored.projects[block_position]
        _set_heading(paragraphs[start], project.title, project.url)
        bullet_paragraphs = [paragraphs[i] for i in range(start + 1, end) if paragraphs[i].text.strip()]
        while len(project.bullets) > len(bullet_paragraphs):
            bullet_paragraphs.append(_clone_paragraph_after(bullet_paragraphs[-1]))
        for i in range(len(bullet_paragraphs) - 1, -1, -1):
            if i >= len(project.bullets):
                _delete_paragraph(bullet_paragraphs[i])
            else:
                _set_paragraph_text(bullet_paragraphs[i], project.bullets[i])

    # Re-map after project deletions, then update only the four technical rows.
    mapping_after = map_template(document)
    for label in SKILL_LABELS[:-1]:
        values = tailored.skills.get(label)
        if values is not None:
            _set_paragraph_text(document.paragraphs[mapping_after.skill_indexes[label]], ", ".join(values), label=label)
    if before != protected_digest(document, map_template(document)):
        raise UnsupportedTemplate("A protected part of the document changed during editing.")
    document.save(output)


def valid_external_url(url: str) -> bool:
    parsed = urlparse(url)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
